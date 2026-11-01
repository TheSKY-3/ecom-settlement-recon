import json

import openpyxl
import pytest
from fastapi.testclient import TestClient

from src.agent import chat, suggestions
from src.workbench import export as export_mod
from src.workbench.app import app
from tests.test_workbench_data import env  # noqa: F401  复用数据夹具


FINDINGS = [
    {"finding_id": "F000001", "order_id": "o1", "diff_type": "单边账",
     "variant": "删账单", "target_side": "bills", "confidence": 1.0,
     "description": "账单缺失：有明细有支付但无账单", "verified": True},
    {"finding_id": "F000002", "order_id": "o2", "diff_type": "金额不符",
     "variant": "费率错", "target_side": "bills", "confidence": 0.9,
     "description": "佣金应为 0.20", "verified": True},
    {"finding_id": "F000003", "order_id": "o3", "diff_type": "退款遗漏",
     "variant": "该退未退", "target_side": "refunds", "confidence": 1.0,
     "description": "取消单已付款但无退款记录", "verified": True},
    {"finding_id": "F000004", "order_id": "o2", "diff_type": "单边账",
     "variant": "删账单", "target_side": "bills", "confidence": 1.0,
     "description": "账单缺失：有明细有支付但无账单", "verified": True},
    {"finding_id": "F000005", "order_id": "o9", "diff_type": "金额不符",
     "variant": "自然差异", "target_side": "bills", "confidence": 0.9,
     "description": "金额与单据不符（无标准答案）", "verified": False},
]


@pytest.fixture()
def export_env(env, tmp_path, monkeypatch):  # noqa: F811
    engine = tmp_path / "engine"
    engine.mkdir()
    (engine / "findings.json").write_text(
        json.dumps({"findings": FINDINGS}, ensure_ascii=False), encoding="utf-8")
    (engine / "score.json").write_text(json.dumps({
        "run_id": "TESTRUN", "gt_total": 4, "finding_total": 5,
        "tp": 4, "fn": 0, "fp": 1, "fp_natural": 1, "fp_other": 0,
        "recall": 1.0, "precision": 0.8, "precision_other": 1.0,
    }), encoding="utf-8")
    monkeypatch.setenv("RECON_ENGINE_DIR", str(engine))
    monkeypatch.setenv("RECON_EXPORT_DIR", str(tmp_path / "exports"))
    export_mod.cleanup()
    yield env
    export_mod.cleanup()


def _read(path):
    return openpyxl.load_workbook(path)


def test_both_export_has_two_sheets_and_pairs(export_env):
    meta = export_mod.create_export({"source": "both"})
    assert meta["rows"] == 5 and meta["orders"] == 4
    assert [s["name"] for s in meta["sheets"]] == ["对照明细", "订单汇总"]
    assert meta["filters"] == {"source": "both"}
    path = export_mod.export_path(meta["token"])
    assert path.is_file() and path.suffix == ".xlsx"

    wb = _read(path)
    assert wb.sheetnames == ["对照明细", "订单汇总"]
    ws = wb["对照明细"]
    assert [c.value for c in ws[1]] == export_mod.DETAIL_HEADERS
    assert ws.max_row == 6 and ws.max_column == len(export_mod.DETAIL_HEADERS)
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref.startswith("A1:")
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    matched = [r for r in rows if r[6] == "一致"]
    assert len(matched) == 4
    solo = [r for r in rows if r[6].startswith("仅引擎")]
    assert len(solo) == 1 and solo[0][0] is None and solo[0][1] == "F000005"
    # 金额是数值、置信度是数值，订单号保持文本（涉案金额在第 12 列，支付额第 16 列）
    pay_cell = ws.cell(row=2, column=16)
    assert pay_cell.value == 15.0 and pay_cell.number_format == "#,##0.00"
    assert ws.cell(row=2, column=10).value == 1.0
    assert ws.cell(row=2, column=12).number_format == "#,##0.00"
    assert ws.cell(row=2, column=3).number_format == "@"

    book = wb["订单汇总"]
    assert [c.value for c in book[1]] == export_mod.SUMMARY_HEADERS
    assert book.max_row == 5
    summary = {r[0]: r for r in book.iter_rows(min_row=2, values_only=True)}
    assert summary["o2"][1] == 2 and summary["o2"][4] == 2
    assert isinstance(summary["o2"][3], (int, float))  # 涉案金额为订单内合计
    assert summary["o9"][4] == 0 and summary["o9"][5] == 1


def test_gt_export_matches_engine_but_only_gt_rows(export_env):
    meta = export_mod.create_export({"source": "gt"})
    assert meta["rows"] == 4 and meta["orders"] == 3
    ws = _read(export_mod.export_path(meta["token"]))["对照明细"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert {r[0] for r in rows} == {"GT000001", "GT000002", "GT000003", "GT000004"}
    assert all(r[6] == "一致" for r in rows)


def test_engine_only_and_natural_filters(export_env):
    natural = export_mod.create_export({"source": "engine", "natural_only": True})
    assert natural["rows"] == 1 and natural["orders"] == 1
    assert natural["filters"]["natural_only"] is True
    conf = export_mod.create_export({"source": "engine", "min_confidence": 1.0})
    assert conf["rows"] == 3
    unverified = export_mod.create_export({"source": "engine", "verified": False})
    assert unverified["rows"] == 1
    one_type = export_mod.create_export({"diff_type": "单边账"})
    assert one_type["rows"] == 2


def test_filters_errors(export_env):
    with pytest.raises(export_mod.ExportError) as e:
        export_mod.create_export({"source": "gt", "min_confidence": 0.9})
    assert e.value.code == "INVALID_ARGS"
    with pytest.raises(export_mod.ExportError) as e:
        export_mod.create_export({"diff_type": "退款遗漏", "q": "没有这个订单"})
    assert e.value.code == "NO_ROWS"


def test_workbook_has_no_formula_cells(export_env):
    meta = export_mod.create_export({"source": "gt"})
    wb = _read(export_mod.export_path(meta["token"]))
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    assert cell.data_type != "f", cell.value


def test_cell_value_guards_formula_and_control_chars():
    assert export_mod._cell_value("=1+1") == "=1+1"
    assert export_mod._cell_value("a\x00b\nc") == "ab\nc"
    assert export_mod._cell_value(True) == "是"
    assert export_mod._cell_value(None) is None
    assert export_mod._cell_value("x" * 600).endswith("…")


def test_token_resolve_and_expiry(export_env, monkeypatch):
    meta = export_mod.create_export({"diff_type": "单边账"})
    assert export_mod.resolve(meta["token"])["rows"] == 2
    assert export_mod.resolve("nope") is None
    monkeypatch.setattr(export_mod, "EXPORT_TTL_S", -1)
    assert export_mod.resolve(meta["token"]) is None
    assert not export_mod.export_path(meta["token"])


def test_api_export_and_tables(export_env):
    client = TestClient(app)
    tables = client.get("/api/tables").json()
    assert len(tables["tables"]) == 8
    assert tables["gt_total"] == 4
    assert tables["engine_run_id"] == "TESTRUN"
    refunds = next(t for t in tables["tables"] if t["key"] == "dirty_refunds")
    assert "refund_reason" in refunds["columns"]

    r = client.post("/api/exports", json={"source": "both", "diff_type": "单边账"})
    assert r.status_code == 200
    body = r.json()
    assert body["rows"] == 2 and body["url"].startswith("/api/exports/")
    got = client.get(body["url"])
    assert got.status_code == 200
    assert got.content[:2] == b"PK"
    assert "spreadsheetml" in got.headers["content-type"]

    bad = client.post("/api/exports", json={"source": "gt", "min_confidence": 0.5})
    assert bad.status_code == 400 and "INVALID_ARGS" in bad.json()["detail"]
    assert client.get("/api/exports/deadbeef").status_code == 404
    assert client.get("/api/chat/presets").json()["presets"][0]["text"]


def test_chat_export_tool_inherits_last_filter(export_env):
    tool = chat.ChatTools()
    session = chat._Session("t")
    out, line = tool.execute("search_diffs", {"source": "gt", "diff_type": "单边账"},
                             session=session)
    assert out["total"] == 2
    meta, line = tool.execute("export_diffs", {}, session=session)
    assert meta["rows"] == 2 and meta["filters"]["diff_type"] == "单边账"
    assert "导出对照表" in line


def test_chat_export_tool_errors_are_facts(export_env):
    tool = chat.ChatTools()
    bad, line = tool.execute("export_diffs", {"source": "gt", "natural_only": True})
    assert bad["error"] == "INVALID_ARGS"
    empty, _ = tool.execute("export_diffs", {"q": "zzz-no-order"})
    assert empty["error"] == "NO_ROWS"


def test_presets_and_dynamic_suggestions_are_executable(export_env):
    tool = chat.ChatTools()
    for preset in suggestions.PRESETS:
        for call in preset["tools"]:
            out, _ = tool.execute(call["name"], call["args"])
            assert not out.get("error"), (preset["text"], call, out)

    facts = [{"kind": "overview", "engine": {"fp_natural": 1}},
             {"kind": "type_stats", "by_type": {"金额不符": {"fp_natural": 1}}}]
    for item in suggestions.build_suggestions(facts):
        for call in item["tools"]:
            out, _ = tool.execute(call["name"], call["args"])
            assert not out.get("error"), (item["text"], call, out)
    exported = suggestions.build_suggestions([{"kind": "export", "url": "/x", "rows": 1}])
    assert exported[0]["text"] == "看看这批里的自然差异样例"


def test_suggestions_tolerate_list_by_type(export_env):
    """get_type_stats 的 by_type 是列表；建议不能按 dict 硬解，更不能抛错。"""
    facts = [{"kind": "type_stats", "by_type": [
        {"diff_type": "金额不符", "fp_natural": 3},
        {"diff_type": "单边账", "fp_natural": 1},
    ]}]
    out = suggestions.build_suggestions(facts)
    assert out[0]["text"] == "看金额不符的 3 条自然差异样例"

    result = chat._decorate(
        chat.ChatTurnResult(session_id="t", status="ready", source="llm", answer="a"),
        [{"kind": "type_stats", "by_type": object()}])
    assert result.suggestions == []
